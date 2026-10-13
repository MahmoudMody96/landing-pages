#!/usr/bin/env python3
"""Add live hero-image columns to the Excel sheet (reads original, writes v2)."""
import json, os, re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = r"D:\MAHMOUD\projects\landing-pages"
DATA = os.path.join(ROOT, "data")
PAGES = os.path.join(ROOT, "src", "cities", "sa-riyadh")

SRC = os.path.join(DATA, "riyadh_businesses_30.xlsx")
OUT = os.path.join(DATA, "riyadh_businesses_30_v2.xlsx")

sel = json.load(open(os.path.join(DATA, "selected_30.json"), encoding="utf-8"))

NICHE_BY_SLUG = [
    ("carve-coffee-bar", "coffee"), ("ik-coffee-downtown", "coffee"),
    ("harmonic-barbershop", "barbershop"), ("twos-barbershop", "barbershop"),
    ("باب-الاناقة", "barbershop"), ("vog-salon", "beauty_salon"),
    ("stellar-beauty-salon", "nail_salon"),
    ("fitness-star-gym", "gym"), ("j-f-revolution", "gym"),
    ("almultaka-ladies-center-fitness", "gym"),
    ("future-of-trust-car-workshop-shafi", "auto_repair"),
    ("car-repairs", "auto_repair"),
]

NICHE_AR = {
    "coffee": "كوفي شوب", "barbershop": "حلاقة رجالية",
    "beauty_salon": "صالون تجميل", "nail_salon": "صالون أظافر",
    "gym": "جيم / لياقة", "auto_repair": "ورشة سيارات",
}


def slugify(n):
    s = n.lower().strip()
    s = re.sub(r"[^\w\s-]", "", s)
    s = re.sub(r"[\s_]+", "-", s)
    return re.sub(r"-+", "-", s).strip("-")[:60]


def hero_of(name):
    slug = slugify(name)
    body = open(os.path.join(PAGES, slug + ".html"), encoding="utf-8").read()
    m = re.search(r'background-image:\s*url\(([^)]+)\)', body)
    url = m.group(1).strip("'\"") if m else ""
    if "pages.aidy.site/hero/" in url:
        for pref, n in NICHE_BY_SLUG:
            if slug.startswith(pref):
                return ("stock", n, url)
    return ("real", "", url)


wb = load_workbook(SRC)
ws = wb["بزنسات الرياض 30"]
ws.insert_cols(9, 2)

thin = Side(style="thin", color="B4C6E7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hfill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
hfont = Font(name="Arial", size=11, bold=True, color="FFFFFF")

for col, label in ((9, "صورة الهيرو"), (10, "مصدر صورة الهيرو")):
    c = ws.cell(row=3, column=col, value=label)
    c.font = hfont
    c.fill = hfill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border

ws.column_dimensions[get_column_letter(9)].width = 46
ws.column_dimensions[get_column_letter(10)].width = 22

stats = {"real": 0, "stock": 0}
for i, b in enumerate(sel, 1):
    kind, niche, url = hero_of(b["name"])
    stats[kind] += 1
    src = "صورة المحل الحقيقية (Google)" if kind == "real" else f"هيرو احترافي — {NICHE_AR.get(niche, niche)}"

    c1 = ws.cell(row=i + 3, column=9, value=url)
    c1.font = Font(name="Arial", size=9, color="0563C1", underline="single")
    c1.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c1.border = border
    c1.hyperlink = url

    c2 = ws.cell(row=i + 3, column=10, value=src)
    c2.font = Font(name="Arial", size=9, bold=(kind == "stock"),
                   color="C55A11" if kind == "stock" else "548235")
    c2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c2.border = border

    if i % 2 == 0:
        for col in (9, 10):
            ws.cell(row=i + 3, column=col).fill = PatternFill(
                start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

ws.auto_filter.ref = "A3:K" + str(len(sel) + 3)

ws2 = wb["ملخص"]
r = ws2.max_row + 2
for k, v in [
    ("صور الهيرو", ""),
    ("صورة المحل الحقيقية (Google)", stats["real"]),
    ("هيرو احترافي (4K) لنفس المجال", stats["stock"]),
    ("", ""),
    ("صور الهيرو الاحترافية", "https://pages.aidy.site/hero/"),
]:
    a = ws2.cell(row=r, column=1, value=k)
    bcell = ws2.cell(row=r, column=2, value=v)
    if k and not v:
        a.font = Font(name="Arial", size=11, bold=True, color="1F3864")
    else:
        a.font = Font(name="Arial", size=10)
    bcell.font = Font(name="Arial", size=10, bold=True)
    a.alignment = Alignment(horizontal="right", vertical="center")
    bcell.alignment = Alignment(horizontal="center", vertical="center")
    if isinstance(v, str) and v.startswith("http"):
        bcell.hyperlink = v
        bcell.font = Font(name="Arial", size=10, color="0563C1", underline="single")
    r += 1

wb.save(OUT)
print(f"saved: {OUT}")
print(f"hero -> real (Google): {stats['real']} | premium 4K: {stats['stock']}")
